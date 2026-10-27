import csv
import io
import json
import re
import zipfile
from decimal import Decimal
from datetime import date, datetime, timedelta
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.config.settings import get_settings
from app.models.workflow import Client, ClientEmail, Company, Employee, GeneratedFile, JobRun, Payroll, ReportJob, SampleSale, UploadedQuery
from app.repositories.workflow_repository import WorkflowRepository
from app.schemas.workflow import EMAIL_PATTERN, JobCreate

APPROVED_PROCEDURES = {
    "mapped_employee_payroll_report": "Mapped Employee Payroll Report",
    "employee_payroll_report": "Employee Payroll Report",
    "monthly_payroll_summary": "Monthly Payroll Summary",
    "department_payroll_report": "Department-wise Payroll Report",
    "employee_salary_breakdown": "Employee Salary Breakdown",
    "company_payroll_report": "Company Payroll Report",
}
MAPPING_FIELDS = (
    "employee_code", "employee_name", "department", "designation", "payroll_month",
    "basic_salary", "allowances", "deductions", "net_salary", "payment_status",
    "company_name", "employee_count", "net_salary_total", "paydate",
)
MAPPING_PROCEDURE = "mapped_employee_payroll_report"
LEGACY_PROCEDURES = {"sales_by_company": "dbo.usp_SalesByCompany", "monthly_sales_summary": "dbo.usp_MonthlySalesSummary", "product_sales_report": "dbo.usp_ProductSalesReport"}
APPROVED_TXT_QUERY = "SELECT e.employee_code AS employee_id, e.employee_name, e.department, e.designation, p.payroll_month, p.basic_salary, p.allowances, p.deductions, p.net_salary, p.payment_status FROM employees e JOIN payroll p ON p.employee_id = e.id WHERE e.company_id = :company_id AND p.payroll_month = (SELECT MAX(p2.payroll_month) FROM payroll p2 JOIN employees e2 ON e2.id = p2.employee_id WHERE e2.company_id = :company_id) ORDER BY e.employee_code"


class WorkflowService:
    def __init__(self, repository: WorkflowRepository | None = None): self.repository = repository or WorkflowRepository()

    def seed(self, db: Session) -> None:
        companies = db.query(Company).order_by(Company.id).all()
        if not companies:
            companies = [Company(company_name=name) for name in ("Company A", "Company B", "Company C")]
            db.add_all(companies); db.flush()
            # Keep the original sample rows intact for existing historical jobs.
            now = datetime.now()
            for company in companies:
                for index, product in enumerate(("Widget", "Gadget", "Service"), start=1):
                    db.add(SampleSale(company_id=company.id, product_name=product, quantity=index * 4, amount=float(index * 125.5), sale_date=now))
            db.commit()
        self.seed_payroll(db, companies)
        self.seed_clients(db, companies)

    def seed_clients(self, db: Session, companies: list[Company]) -> None:
        if db.query(Client).count(): return
        for company in companies:
            for index in range(1, 4):
                client = Client(company_id=company.id, client_code=f"CLIENT{company.id:02d}{index:03d}", client_name=f"{company.company_name} Client {index}")
                db.add(client); db.flush()
                db.add_all([ClientEmail(client_id=client.id, email=f"finance{index}@company{company.id}.example"), ClientEmail(client_id=client.id, email=f"payroll{index}@company{company.id}.example")])
        db.commit()

    def seed_payroll(self, db: Session, companies: list[Company]) -> None:
        if db.query(Employee).count(): return
        departments = (("Finance", "Payroll Analyst"), ("Operations", "Operations Specialist"), ("Technology", "Systems Analyst"), ("People", "People Coordinator"), ("Sales", "Account Specialist"))
        names = ("Avery Morgan", "Jordan Ellis", "Riley Hart", "Casey Rowan", "Taylor Quinn", "Morgan Reed", "Cameron Blake", "Jamie Vale", "Drew Parker", "Skyler Lane", "Emery Stone", "Robin Avery")
        current_month = datetime(datetime.now().year, datetime.now().month, 1)
        months = [(current_month - timedelta(days=32 * offset)).replace(day=1) for offset in (2, 1, 0)]
        for company_index, company in enumerate(companies, start=1):
            for index in range(12):
                department, designation = departments[index % len(departments)]
                employee = Employee(employee_code=f"D{company_index:02d}-{index + 1:04d}", employee_name=names[(company_index * 12 + index) % len(names)], company_id=company.id, department=department, designation=designation, is_active=True)
                db.add(employee); db.flush()
                for month_index, month in enumerate(months):
                    basic = Decimal(4200 + company_index * 275 + index * 85 + month_index * 20)
                    allowance = (basic * Decimal("0.12")).quantize(Decimal("0.01")); deduction = (basic * Decimal("0.07")).quantize(Decimal("0.01"))
                    db.add(Payroll(employee_id=employee.id, payroll_month=month, basic_salary=basic, allowances=allowance, deductions=deduction, net_salary=basic + allowance - deduction, payment_status="Paid" if month_index < 2 else "Processing"))
        db.commit()

    def create_upload(self, db: Session, original_name: str, content: bytes) -> UploadedQuery:
        settings = get_settings()
        extension = Path(original_name).suffix.lower()
        if extension not in {".txt", ".xlsx", ".xlsm"} or len(content) > settings.max_upload_bytes:
            raise ValueError("Upload must be a TXT or Excel file within the configured size limit.")
        if extension in {".xlsx", ".xlsm"}:
            return self._create_excel_column_upload(db, original_name, content, settings)
        try:
            query = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("TXT uploads must use UTF-8 encoding.") from exc
        # Keep the original approved TXT-query workflow available for existing users.
        if self._normalize(query) == self._normalize(APPROVED_TXT_QUERY):
            folder = Path(settings.uploads_dir); folder.mkdir(parents=True, exist_ok=True)
            stored_name = f"{uuid4().hex}.txt"; (folder / stored_name).write_text(query.strip(), encoding="utf-8")
            record = UploadedQuery(original_name=original_name, stored_name=stored_name, approved_key="employee_payroll_report", column_names=None)
            db.add(record); db.commit(); db.refresh(record); return record
        names = [line.strip() for line in query.splitlines()]
        if not names or any(not name for name in names):
            raise ValueError("TXT uploads must contain one non-blank column name per line.")
        if len(names) > len(MAPPING_FIELDS):
            raise ValueError("TXT uploads contain too many columns for the approved payroll report.")
        if any(len(name) > 160 or any(ord(char) < 32 for char in name) for name in names):
            raise ValueError("TXT column names contain invalid characters or are too long.")
        if len({name.casefold() for name in names}) != len(names):
            raise ValueError("TXT column names must be unique.")
        folder = Path(settings.uploads_dir); folder.mkdir(parents=True, exist_ok=True)
        stored_name = f"{uuid4().hex}.txt"; (folder / stored_name).write_text("\n".join(names) + "\n", encoding="utf-8")
        record = UploadedQuery(original_name=original_name, stored_name=stored_name, approved_key=MAPPING_PROCEDURE, column_names=json.dumps(names, ensure_ascii=False))
        db.add(record); db.commit(); db.refresh(record); return record

    def _create_excel_column_upload(self, db: Session, original_name: str, content: bytes, settings) -> UploadedQuery:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows = workbook.active.iter_rows(values_only=True)
            header = [str(value or "").strip().casefold() for value in next(rows, ())]
            if "record type" not in header or "column name" not in header:
                raise ValueError("Excel file must contain 'record type' and 'column name' columns.")
            name_index = header.index("column name")
            record_type_index = header.index("record type")
            columns = []
            for row in rows:
                if len(row) <= name_index or row[name_index] is None:
                    continue
                name = str(row[name_index]).strip()
                if not name:
                    continue
                raw_record_type = row[record_type_index] if len(row) > record_type_index else None
                record_type = str(raw_record_type or "D").strip().upper()
                if record_type not in {"H", "D", "T"}:
                    raise ValueError("Excel record type must be H, D, or T.")
                columns.append({"column_name": name, "record_type": record_type})
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Excel file could not be read. Upload a valid .xlsx workbook.") from exc
        finally:
            try: workbook.close()
            except UnboundLocalError: pass
        self._validate_mapping_names(columns, "Excel")
        folder = Path(settings.uploads_dir); folder.mkdir(parents=True, exist_ok=True)
        stored_name = f"{uuid4().hex}{Path(original_name).suffix.lower()}"; (folder / stored_name).write_bytes(content)
        record = UploadedQuery(original_name=original_name, stored_name=stored_name, approved_key=MAPPING_PROCEDURE, column_names=json.dumps(columns, ensure_ascii=False))
        db.add(record); db.commit(); db.refresh(record); return record

    @staticmethod
    def _validate_mapping_names(columns: list[dict], source_label: str) -> None:
        if not columns: raise ValueError(f"{source_label} must contain at least one column name.")
        if len(columns) > len(MAPPING_FIELDS): raise ValueError(f"{source_label} contains too many columns for the approved payroll report.")
        if any(len(item["column_name"]) > 160 or any(ord(char) < 32 for char in item["column_name"]) for item in columns): raise ValueError(f"{source_label} column names contain invalid characters or are too long.")
        keys = [(item.get("record_type", "D").upper(), item["column_name"].casefold()) for item in columns]
        if len(set(keys)) != len(keys): raise ValueError(f"{source_label} column names must be unique within each record type.")

    def _job_values(self, db: Session, payload: JobCreate) -> dict:
        if db.get(Company, payload.company_id) is None: raise ValueError("Company not found.")
        client_selections = {}
        for raw_client_id, email_ids in payload.client_selections.items():
            try: client_id = int(raw_client_id)
            except (TypeError, ValueError) as exc: raise ValueError("Client selection is invalid.") from exc
            if db.get(Company, client_id) is None: raise ValueError("Every selected client must be a valid company ID.")
            normalized_emails = []
            seen_emails = set()
            for raw_email in email_ids:
                email = str(raw_email).strip()
                # Numeric values are accepted only for compatibility with older
                # jobs that stored ClientEmail IDs. New jobs store manual emails.
                if not email.isdigit() and not EMAIL_PATTERN.fullmatch(email):
                    raise ValueError(f"Invalid email address: {email or '(empty)'}")
                key = email.casefold()
                if key in seen_emails:
                    raise ValueError("Company email recipients must be unique.")
                seen_emails.add(key); normalized_emails.append(email)
            client_selections[str(client_id)] = normalized_emails
        if payload.source_type == "procedure" and payload.source_key not in APPROVED_PROCEDURES: raise ValueError("Choose an approved stored procedure.")
        if payload.source_type == "txt" and not payload.uploaded_query_id: raise ValueError("Choose an approved TXT upload.")
        mappings = payload.column_mappings or []
        if payload.source_key == MAPPING_PROCEDURE:
            upload = self.repository.uploaded(db, payload.uploaded_query_id) if payload.uploaded_query_id else None
            if not upload or upload.approved_key != MAPPING_PROCEDURE:
                raise ValueError("Upload a valid TXT column list for the mapped payroll procedure.")
            uploaded_columns = json.loads(upload.column_names or "[]")
            expected_names = [item.get("column_name", "") if isinstance(item, dict) else item for item in uploaded_columns]
            if [item.column_name for item in mappings] != expected_names:
                raise ValueError("Every upload column must be mapped in its original order.")
            if any(field and field not in MAPPING_FIELDS for item in mappings for field in (item.db_column_name if isinstance(item.db_column_name, list) else [item.db_column_name])):
                raise ValueError("Choose fields from the approved payroll field list.")
        elif mappings or payload.uploaded_query_id:
            raise ValueError("TXT mappings are supported only by the mapped payroll procedure.")
        values = payload.model_dump(exclude={"column_mappings", "schedule_type", "weekly_day", "monthly_day", "sort_field", "sort_direction", "filter_field", "filter_operator", "filter_value", "sort_rules", "filter_rules"})
        values["recipients"] = json.dumps(payload.recipients, ensure_ascii=False) if payload.recipients else None
        values["client_selections"] = json.dumps(client_selections, ensure_ascii=False) if client_selections else None
        values["column_mappings"] = json.dumps([item.model_dump(exclude_none=True) for item in mappings], ensure_ascii=False) if mappings else None
        sort_rules = payload.sort_rules or ([{"field": payload.sort_field, "direction": payload.sort_direction}] if payload.sort_field else [])
        filter_rules = payload.filter_rules or ([{"field": payload.filter_field, "operator": payload.filter_operator, "value": str(payload.filter_value)}] if payload.filter_field and payload.filter_value is not None else [])
        values["report_options"] = json.dumps({"sort_rules": sort_rules, "filter_rules": filter_rules, "sort_field": payload.sort_field, "sort_direction": payload.sort_direction, "filter_field": payload.filter_field, "filter_operator": payload.filter_operator, "filter_value": str(payload.filter_value) if payload.filter_value is not None else None}, ensure_ascii=False)
        recurrence = {"one_time": "once", "every_2_minutes": "every_2_minutes", "hourly_1": "hourly_1", "hourly_2": "hourly_2", "hourly_4": "hourly_4", "daily": "daily", "weekly": "weekly", "monthly": "monthly"}[payload.schedule_type]
        if payload.schedule_type == "weekly" and payload.weekly_day is None: raise ValueError("Choose a day of the week for weekly schedules.")
        if payload.schedule_type == "monthly" and payload.monthly_day is None: raise ValueError("Choose a day of the month for monthly schedules.")
        values["schedule_time"] = payload.schedule_at.strftime("%H:%M"); values["recurrence"] = recurrence; values["day_of_week"] = payload.weekly_day; values["day_of_month"] = payload.monthly_day; values["next_run_at"] = payload.schedule_at
        return values

    def create_job(self, db: Session, payload: JobCreate) -> ReportJob:
        values = self._job_values(db, payload)
        job = ReportJob(**values); db.add(job); db.commit(); db.refresh(job); return job

    def update_job(self, db: Session, job_id: int, payload: JobCreate) -> ReportJob:
        job = self.repository.job(db, job_id)
        if not job: raise ValueError("Job not found.")
        values = self._job_values(db, payload)
        for key, value in values.items():
            if key not in {"id", "created_at"}:
                setattr(job, key, value)
        db.commit(); db.refresh(job); return job

    def deletion_summary(self, db: Session) -> dict[str, int]:
        jobs = db.query(ReportJob).all()
        job_ids = [job.id for job in jobs]
        run_count = db.query(JobRun).filter(JobRun.job_id.in_(job_ids)).count() if job_ids else 0
        file_count = db.query(GeneratedFile).filter(GeneratedFile.job_id.in_(job_ids)).count() if job_ids else 0
        upload_ids = {job.uploaded_query_id for job in jobs if job.uploaded_query_id}
        return {"jobs": len(jobs), "runs": run_count, "generated_files": file_count, "uploads": len(upload_ids)}

    def delete_all_jobs_and_files(self, db: Session) -> dict[str, int]:
        jobs = db.query(ReportJob).all()
        job_ids = [job.id for job in jobs]
        upload_ids = {job.uploaded_query_id for job in jobs if job.uploaded_query_id}
        files = db.query(GeneratedFile).filter(GeneratedFile.job_id.in_(job_ids)).all() if job_ids else []
        runs = db.query(JobRun).filter(JobRun.job_id.in_(job_ids)).all() if job_ids else []
        uploads = db.query(UploadedQuery).filter(UploadedQuery.id.in_(upload_ids)).all() if upload_ids else []
        generated_root = Path(get_settings().generated_reports_dir).resolve()
        upload_root = Path(get_settings().uploads_dir).resolve()
        for record in files:
            path = Path(record.storage_path).resolve()
            if path.is_file() and generated_root in path.parents:
                path.unlink()
        for record in uploads:
            path = (upload_root / Path(record.stored_name).name).resolve()
            if path.is_file() and upload_root in path.parents:
                path.unlink()
        for record in files: db.delete(record)
        for record in runs: db.delete(record)
        for record in jobs: db.delete(record)
        for record in uploads: db.delete(record)
        db.commit()
        return {"jobs": len(jobs), "runs": len(runs), "generated_files": len(files), "uploads": len(uploads)}

    def set_enabled(self, db: Session, job_id: int, enabled: bool) -> ReportJob:
        job = self.repository.job(db, job_id)
        if not job: raise ValueError("Job not found.")
        job.enabled = enabled; db.commit(); db.refresh(job); return job

    def run_job(self, db: Session, job_id: int, retry_run_id: int | None = None) -> JobRun:
        job = self.repository.job(db, job_id)
        if not job: raise ValueError("Job not found.")
        if not job.enabled and retry_run_id is None:
            raise ValueError("Job is disabled. Enable it before running.")
        run = self.repository.run(db, retry_run_id) if retry_run_id else None
        if run and run.status not in {"Failed", "Partially Completed"}: raise ValueError("Only failed runs can be retried.")
        is_new = run is None
        if is_new:
            run = JobRun(job_id=job.id, status="Running", current_step="SQL execution", step_log="[]")
            db.add(run); db.commit(); db.refresh(run)
        if not is_new and run.status == "Running": raise ValueError("This run is already active.")
        run.status = "Running"; run.error_message = None
        try:
            if not run.csv_content and not run.output_manifest:
                if job.export_split_by and job.export_split_by != "file_name":
                    split_outputs = self._execute_split_source(db, job)
                    run.output_manifest = json.dumps([{"name": name, "content": content, "rows": row_count} for name, content, row_count in split_outputs], ensure_ascii=False)
                    run.csv_content = split_outputs[0][1] if split_outputs else ""
                else:
                    rows, headers = self._execute_source(db, job)
                    run.csv_content = self._csv_records(rows, job.delimiter or ",") if headers is None else self._csv(headers, rows, job.delimiter or ",")
                self._log(run, "SQL execution", "Completed")
                run.current_step = "CSV generation"
                self._log(run, "CSV generation", "Completed")
            run.current_step = "File storage"
            generated = self._store_file(db, job, run)
            self._log(run, "File storage", "Completed")
            run.status = "Completed"; run.current_step = "Completed"; run.completed_at = datetime.now(); job.last_run_status = run.status
            db.commit(); db.refresh(run); return run
        except Exception as exc:
            run.status = "Failed" if not run.csv_content else "Partially Completed"; run.error_message = "Report execution failed. Review the run details."
            run.current_step = run.current_step or "Unknown"; self._log(run, run.current_step, "Failed")
            job.last_run_status = run.status; db.commit(); raise ValueError(run.error_message) from exc

    def _execute_source(self, db: Session, job: ReportJob):
        dialect = db.bind.dialect.name
        if job.source_key == MAPPING_PROCEDURE:
            mappings = json.loads(job.column_mappings or "[]")
            # Normalize mappings saved by older versions and the new four-column UI.
            normalized = []
            for item in mappings:
                column_name = item.get("column_name") or item.get("display_name")
                database_field = item.get("db_column_name") or item.get("database_field")
                database_fields = database_field if isinstance(database_field, list) else [database_field]
                normalized.append({**item, "column_name": column_name, "display_name": item.get("display_name") or column_name, "db_column_name": database_field, "database_field": database_field, "record_type": item.get("record_type") or "D", "format": item.get("format") or "Text"})
            mappings = normalized
            selected_company_ids = self._selected_company_ids(job)
            if dialect == "mssql" and not selected_company_ids and not self._requires_application_mapping(mappings):
                result = db.execute(text("EXEC dbo.usp_GenerateMappedPayrollReport @company_id=:company_id, @column_mappings=:column_mappings"), {"company_id": job.company_id, "column_mappings": json.dumps(mappings, ensure_ascii=False)})
            else:
                return self._execute_application_mapping(db, job, mappings)
            rows = result.fetchmany(get_settings().max_report_rows)
            return [tuple(row) for row in rows], list(result.keys())
        if job.source_type == "procedure" and dialect == "mssql":
            names = {"employee_payroll_report": "dbo.usp_EmployeePayrollReport", "monthly_payroll_summary": "dbo.usp_MonthlyPayrollSummary", "department_payroll_report": "dbo.usp_DepartmentPayrollReport", "employee_salary_breakdown": "dbo.usp_EmployeeSalaryBreakdown", "company_payroll_report": "dbo.usp_CompanyPayrollReport", **LEGACY_PROCEDURES}
            procedure_name = names[job.source_key]
            if job.source_key == "company_payroll_report":
                # This legacy procedure is parameterless. Execute it as-is,
                # then apply the selected, verified company IDs to its
                # company_name column in the application layer.
                result = db.execute(text(f"EXEC {procedure_name}"))
                rows = [tuple(row) for row in result.fetchmany(get_settings().max_report_rows)]
                selected_ids = self._selected_company_ids(job) or [job.company_id]
                companies = db.query(Company).filter(Company.id.in_(selected_ids)).all()
                if len(companies) != len(set(selected_ids)) or len({company.company_name for company in companies}) != len(companies):
                    raise ValueError("Selected company identities could not be verified for stored-procedure filtering.")
                allowed_names = {company.company_name for company in companies}
                # usp_CompanyPayrollReport returns RECTYPE at index 1 and the
                # unique company_name at index 2. Keep all H/D/T rows for the
                # selected companies and preserve the procedure's ordering.
                rows = [row for row in rows if len(row) > 2 and row[2] in allowed_names]
                return rows, list(result.keys())
            result = db.execute(text(f"EXEC {procedure_name} @company_id=:company_id"), {"company_id": job.company_id})
        elif job.source_type == "procedure" and job.source_key in APPROVED_PROCEDURES:
            result = db.execute(text(APPROVED_TXT_QUERY), {"company_id": job.company_id})
        elif job.source_type == "txt":
            result = db.execute(text(APPROVED_TXT_QUERY), {"company_id": job.company_id})
        else:
            result = db.execute(text("SELECT product_name, quantity, amount, sale_date FROM sample_sales WHERE company_id = :company_id ORDER BY sale_date DESC"), {"company_id": job.company_id})
        rows = result.fetchmany(get_settings().max_report_rows)
        return [tuple(row) for row in rows], list(result.keys())

    def _store_file(self, db: Session, job: ReportJob, run: JobRun) -> GeneratedFile:
        existing = db.query(GeneratedFile).filter_by(run_id=run.id).first()
        if existing: return existing
        company = db.get(Company, job.company_id)
        folder = Path(get_settings().generated_reports_dir); folder.mkdir(parents=True, exist_ok=True)
        if run.output_manifest:
            manifest = json.loads(run.output_manifest)
            stamp = f"{datetime.now():%Y%m%d_%H%M%S}"
            # Split archives may contain several companies; avoid naming the
            # archive after only the primary company shown on the job record.
            job_slug = re.sub(r"[^A-Za-z0-9_-]+", "_", (job.title or "payroll_export").strip()).strip("_").lower() or "payroll_export"
            archive_name = f"{job_slug}_{job.id}_run_{run.id}_{stamp}.zip"
            archive_path = folder / archive_name
            total_rows = 0
            with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for item in manifest:
                    archive.writestr(item["name"], item["content"])
                    total_rows += item.get("rows", 0)
            generated = GeneratedFile(run_id=run.id, job_id=job.id, company_id=job.company_id, file_name=archive_name, storage_path=str(archive_path.resolve()), row_count=total_rows, file_size=archive_path.stat().st_size)
            db.add(generated); return generated
        filename = f"{company.company_name.lower().replace(' ', '_')}_{job.id}_run_{run.id}_{datetime.now():%Y%m%d_%H%M%S}.csv"
        path = folder / filename; path.write_text(run.csv_content or "", encoding="utf-8", newline="")
        row_count = self._generated_detail_row_count(job, run.csv_content or "")
        generated = GeneratedFile(run_id=run.id, job_id=job.id, company_id=job.company_id, file_name=filename, storage_path=str(path.resolve()), row_count=row_count, file_size=path.stat().st_size)
        db.add(generated); return generated

    def _generated_detail_row_count(self, job: ReportJob, content: str) -> int:
        """Return D-record count for H/D/T output, legacy header-row count otherwise."""
        mappings = json.loads(job.column_mappings or "[]") if job.column_mappings else []
        record_types = [str(item.get("record_type") or "D").upper() for item in mappings if isinstance(item, dict)]
        if "H" not in record_types and "T" not in record_types:
            return max(0, len(content.splitlines()) - 1)
        delimiter = job.delimiter if job.delimiter in {",", "|", ";", "\t"} else ","
        physical_rows = list(csv.reader(io.StringIO(content), delimiter=delimiter))
        return max(0, len(physical_rows) - (1 if "H" in record_types else 0) - (1 if "T" in record_types else 0))

    def _execute_split_source(self, db: Session, job: ReportJob):
        if job.source_key != MAPPING_PROCEDURE:
            raise ValueError("Export splitting is supported for the mapped payroll report only.")
        raw_rows, mappings = self._query_application_rows(db, job)
        groups = {}
        for row in raw_rows:
            # Pay Date exports are kept separate per selected company and pay
            # date so a company's report is never mixed with another company's
            # records that happen to share the same payroll month.
            key = row["__client_id"] if job.export_split_by in {"client_id", "file_client"} else (row["__client_id"], row["__pay_date"])
            groups.setdefault(key, []).append(row)
        outputs = []
        for key, rows in groups.items():
            if job.export_split_by in {"client_id", "file_client"}:
                safe_key = re.sub(r"[^A-Za-z0-9_-]+", "_", str(key))
                filename = f"client_{safe_key}_{job.id}.csv"
            else:
                client_id, pay_date = key
                safe_date = re.sub(r"[^A-Za-z0-9_-]+", "_", str(pay_date))
                filename = f"client_{client_id}_paydate_{safe_date}_{job.id}.csv"
            if any(item["record_type"] in {"H", "T"} for item in mappings):
                content = self._csv_records(self._build_htd_records(mappings, rows), job.delimiter or ",")
            else:
                headers = [item["display_name"] for item in mappings]
                detail_rows = [tuple(self._mapping_value(row, item) for item in mappings) for row in rows]
                content = self._csv(headers, detail_rows, job.delimiter or ",")
            outputs.append((filename, content, len(rows)))
        return outputs

    @staticmethod
    def _requires_application_mapping(mappings):
        procedure_fields = {"employee_code", "employee_name", "department", "designation", "payroll_month", "basic_salary", "allowances", "deductions", "net_salary", "payment_status"}
        return any(item.get("record_type", "D") != "D" or any(field not in procedure_fields for field in (item.get("database_field") if isinstance(item.get("database_field"), list) else [item.get("database_field") or item.get("db_column_name")])) or len(item.get("database_field") if isinstance(item.get("database_field"), list) else [item.get("database_field") or item.get("db_column_name")]) > 1 for item in mappings)

    def _query_application_rows(self, db: Session, job: ReportJob):
        mappings = json.loads(job.column_mappings or "[]")
        normalized = []
        for item in mappings:
            fields = item.get("database_field") or item.get("db_column_name")
            fields = fields if isinstance(fields, list) else [fields]
            fields = [field for field in fields if field]
            if any(field not in MAPPING_FIELDS for field in fields): raise ValueError("Choose fields from the approved payroll field list.")
            normalized.append({"display_name": item.get("display_name") or item.get("column_name"), "database_field": fields, "record_type": item.get("record_type") or "D", "format": item.get("format") or "Text", "format_pattern": item.get("format_pattern")})
        expressions = {
            "employee_id": "e.id", "employee_code": "e.employee_code", "employee_name": "e.employee_name", "department": "e.department", "designation": "e.designation",
            "payroll_month": "p.payroll_month", "paydate": "e.paydate", "basic_salary": "p.basic_salary", "allowances": "p.allowances", "deductions": "p.deductions",
            "net_salary": "p.net_salary", "payment_status": "p.payment_status", "company_name": "c.company_name", "company_id": "e.company_id", "is_active": "e.is_active",
        }
        selected_fields = {field for item in normalized for field in item["database_field"] if field in expressions}
        if "net_salary_total" in {field for item in normalized for field in item["database_field"]}:
            selected_fields.add("net_salary")
        # Pay Date splitting uses the employee-level test field when present.
        # Payroll-period reporting remains available separately through the
        # payroll_month mapping field.
        select_parts = ["e.company_id AS __client_id", "e.paydate AS __pay_date"]
        for field in selected_fields:
            select_parts.append(f'{expressions[field]} AS "__{field}"')
        ids = self._selected_company_ids(job) or [job.company_id]
        try: options = json.loads(job.report_options or "{}")
        except (TypeError, json.JSONDecodeError): options = {}
        sortable = {field: expressions[field] for field in expressions}
        numeric = {field: expressions[field] for field in ("basic_salary", "allowances", "deductions", "net_salary", "employee_id", "company_id", "is_active")}
        text_fields = {field: expressions[field] for field in ("employee_code", "employee_name", "department", "designation", "payment_status", "company_name")}
        dates = {field: expressions[field] for field in ("payroll_month", "paydate")}
        clauses = [f"e.company_id IN ({','.join(str(int(value)) for value in ids)})"]; params = {}
        operators = {"eq": "=", "neq": "!=", "gt": ">", "gte": ">=", "lt": "<", "lte": "<="}
        filter_rules = options.get("filter_rules") or ([] if not options.get("filter_field") else [{"field": options.get("filter_field"), "operator": options.get("filter_operator"), "value": options.get("filter_value")}])
        for index, rule in enumerate(filter_rules):
            field, operator, value = rule.get("field"), rule.get("operator"), rule.get("value")
            field_map = numeric.get(field) or text_fields.get(field) or dates.get(field)
            if not field_map:
                continue
            if operator in {"is_null", "is_not_null", "is_empty", "is_not_empty"}:
                null_clause = f"{field_map} IS NULL" if operator in {"is_null", "is_empty"} else f"{field_map} IS NOT NULL"
                if operator in {"is_empty", "is_not_empty"} and field in text_fields:
                    null_clause = f"COALESCE({field_map}, '') = ''" if operator == "is_empty" else f"COALESCE({field_map}, '') <> ''"
                clauses.append(null_clause)
                continue
            if field_map and value not in (None, "", []) and operator in (set(operators) | {"contains", "not_contains", "starts_with", "ends_with", "in"}):
                parameter = f"filter_value_{index}"
                if field in text_fields and operator in {"contains", "starts_with", "ends_with"}:
                    pattern = {"contains": f"%{value}%", "starts_with": f"{value}%", "ends_with": f"%{value}"}[operator]
                    clauses.append(f"{field_map} LIKE :{parameter}"); params[parameter] = pattern
                elif field in text_fields and operator == "not_contains":
                    clauses.append(f"{field_map} NOT LIKE :{parameter}"); params[parameter] = f"%{value}%"
                elif operator == "in":
                    values = value if isinstance(value, list) else [item.strip() for item in str(value).split(",") if item.strip()]
                    placeholders = []
                    for item_index, item in enumerate(values):
                        item_parameter = f"{parameter}_{item_index}"; placeholders.append(f":{item_parameter}"); params[item_parameter] = item
                    if placeholders: clauses.append(f"{field_map} IN ({', '.join(placeholders)})")
                elif operator not in {"contains", "starts_with", "ends_with"}:
                    clauses.append(f"{field_map} {operators[operator]} :{parameter}"); params[parameter] = value
        sort_rules = options.get("sort_rules") or ([] if not options.get("sort_field") else [{"field": options.get("sort_field"), "direction": options.get("sort_direction")}])
        order_parts = [f"{sortable[rule['field']]} {'DESC' if rule.get('direction') == 'desc' else 'ASC'}" for rule in sort_rules if rule.get("field") in sortable]
        order_parts = order_parts or ["e.company_id ASC", "p.payroll_month ASC", "e.employee_code ASC"]
        result = db.execute(text(f"SELECT {', '.join(select_parts)} FROM employees e JOIN companies c ON c.id=e.company_id JOIN payroll p ON p.employee_id=e.id WHERE {' AND '.join(clauses)} ORDER BY {', '.join(order_parts)}"), params)
        rows = [dict(row._mapping) for row in result.fetchmany(get_settings().max_report_rows)]
        return rows, normalized

    def _execute_application_mapping(self, db: Session, job: ReportJob, mappings):
        rows, normalized = self._query_application_rows(db, job)
        if self._requires_application_mapping(mappings):
            return self._build_htd_records(normalized, rows), None
        headers = [item["display_name"] for item in normalized]
        return [tuple(self._mapping_value(row, item) for item in normalized) for row in rows], headers

    def _build_htd_records(self, mappings, rows):
        if not rows: return []
        records = []
        for record_type in ("H", "D", "T"):
            selected = [item for item in mappings if item["record_type"] == record_type]
            if not selected: continue
            source_rows = rows if record_type == "D" else [rows[0]]
            if record_type == "T": source_rows = [rows[0]]
            for source in source_rows:
                values = []
                for item in selected:
                    value = self._mapping_value(source, item, rows)
                    values.append(self._format_value(value, item["format"], item.get("format_pattern")))
                records.append(values)
        return records

    def _mapping_value(self, row, item, all_rows=None):
        fields = item.get("database_field") or item.get("db_column_name") or []
        fields = fields if isinstance(fields, list) else [fields]
        values = []
        for field in fields:
            if field == "employee_count": value = len(all_rows or [])
            elif field == "net_salary_total": value = sum((Decimal(str(r.get("__net_salary") or 0)) for r in (all_rows or [row])), Decimal("0"))
            elif field == "company_name" and all_rows is not None and len({r.get("__client_id") for r in all_rows}) > 1: value = "Multiple Companies"
            else: value = row.get(f"__{field}")
            if value is not None and str(value) != "": values.append(str(value))
        return " ".join(values)

    def _csv_records(self, records, delimiter):
        delimiter = delimiter if delimiter in {",", "|", ";", "\t"} else ","
        output = io.StringIO(); writer = csv.writer(output, delimiter=delimiter)
        writer.writerows([[self._safe(value) for value in row] for row in records]); return output.getvalue()

    @staticmethod
    def _format_value(value, format_name, format_pattern=None):
        if value is None:
            return str(format_pattern) if format_pattern else ""
        if format_name == "Default":
            text_value = str(value)
            if format_pattern:
                try: return str(format_pattern).replace("{value}", text_value)
                except Exception: return text_value
            return text_value
        if format_name in {"Date", "DateDMY", "DateMDY"} and isinstance(value, (datetime, date)):
            return value.strftime({"Date": "%Y-%m-%d", "DateDMY": "%d-%m-%Y", "DateMDY": "%m/%d/%Y"}[format_name])
        if format_name in {"PhoneDash", "PhoneParens"}:
            digits = re.sub(r"\D", "", str(value))
            if len(digits) == 10:
                return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}" if format_name == "PhoneDash" else f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
        if format_name in {"CurrencyUSD", "CurrencyINR", "CurrencyEUR"}:
            symbols = {"CurrencyUSD": "$", "CurrencyINR": "₹", "CurrencyEUR": "€"}
            try: return f"{symbols[format_name]}{Decimal(str(value)):,.2f}"
            except Exception: return str(value)
        return str(value)

    @staticmethod
    def _selected_company_ids(job: ReportJob) -> list[int]:
        try:
            selections = json.loads(job.client_selections or "{}")
            return sorted({int(value) for value in selections})
        except (TypeError, ValueError, json.JSONDecodeError):
            return []

    @staticmethod
    def _normalize(query: str) -> str: return re.sub(r"\s+", " ", query.strip()).lower()
    @staticmethod
    def _safe(value): return "'" + value if isinstance(value, str) and value[:1] in ("=", "+", "-", "@") else value
    def _csv(self, headers, rows, delimiter=","):
        delimiter = delimiter if delimiter in {",", "|", ";", "\t"} else ","
        output = io.StringIO(); writer = csv.writer(output, delimiter=delimiter); writer.writerow(headers); writer.writerows([[self._safe(v) for v in row] for row in rows]); return output.getvalue()
    @staticmethod
    def _log(run, step, state):
        history = json.loads(run.step_log or "[]"); history.append({"step": step, "status": state, "at": datetime.now().isoformat()}); run.step_log = json.dumps(history)
